"""MARSHALL brake-pad cross-reference index.

The manufacturer's public catalogue is an XLSX download.  Customer searches
must never download or parse that six-megabyte file, so a reviewed snapshot is
compiled into a small read-only SQLite index beforehand.  The index keeps only
active passenger-car brake pads for this W1 adapter.

The adapter is intentionally a candidate: it is invisible unless an explicit
``CP_PILOT_RULES`` entry enables its exact group and cohort.
"""
from __future__ import annotations

import asyncio
import hashlib
import os
import re
import sqlite3
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from ..groups import BRAKE_PADS
from ..normalize import (
    KIND_AFTERMARKET, KIND_OEM, clean_brand, clean_number, number_key, split_brands,
)
from .base import BaseSource, Cross, SourceResult, SourceStatus


CATALOGUE_URL = (
    "https://cars.marshall.parts/wp-content/uploads/2026/03/"
    "marshall-catalog-cars-24.03.2026.xlsx"
)
HOMEPAGE = "https://cars.marshall.parts/shop/"
_REFERENCE = re.compile(r"\s*([^();]+?)\s*\(([^()]+)\)\s*")


@dataclass(frozen=True)
class MarshallReference:
    brand: str
    number: str
    kind: str


@dataclass(frozen=True)
class MarshallProduct:
    article: str
    url: str
    references: tuple[MarshallReference, ...]


def parse_references(value: object, *, kind: str) -> list[MarshallReference]:
    """Parse ``NUMBER(BRAND); …`` cells from MARSHALL's public workbook.

    The source uses this notation in both the original-number and analogue
    columns.  Text that does not have an explicit brand is deliberately
    ignored: a price, a vehicle name or a dangling product name must not become
    a cross-reference.
    """
    text = str(value or "").strip()
    if not text or text == "-":
        return []
    out: list[MarshallReference] = []
    seen: set[tuple[str, str]] = set()
    for chunk in text.split(";"):
        match = _REFERENCE.fullmatch(chunk)
        if match is None:
            continue
        number = clean_number(match.group(1))
        if not number:
            continue
        for raw_brand in split_brands(match.group(2)):
            brand = clean_brand(raw_brand)
            key = brand, number_key(number)
            if not brand or key in seen:
                continue
            seen.add(key)
            out.append(MarshallReference(brand, number, kind))
    return out


def brake_pad_rows(workbook_path: Path) -> list[MarshallProduct]:
    """Read active passenger-car brake pads from the official spreadsheet."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header = {
            str(value or "").strip(): index
            for index, value in enumerate(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        }
        required = {
            "Артикул MARSHALL", "Статус", "Категория", "Оригинальные номера",
            "Аналоги", "Ссылка на страницу товара",
        }
        missing = required - set(header)
        if missing:
            raise ValueError(f"в XLSX MARSHALL нет колонок: {', '.join(sorted(missing))}")
        products: list[MarshallProduct] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            article = clean_number(str(row[header["Артикул MARSHALL"]] or ""))
            status = str(row[header["Статус"]] or "").strip().casefold()
            category = str(row[header["Категория"]] or "").strip().casefold()
            if not article or status != "активный ассортимент" or not category.startswith("тормозные колодки"):
                continue
            url = str(row[header["Ссылка на страницу товара"]] or "").strip()
            references = (
                *parse_references(row[header["Оригинальные номера"]], kind=KIND_OEM),
                *parse_references(row[header["Аналоги"]], kind=KIND_AFTERMARKET),
            )
            if references:
                products.append(MarshallProduct(article, url, tuple(references)))
        return products
    finally:
        workbook.close()


class MarshallIndex:
    def __init__(self, products: list[MarshallProduct]):
        self.products = products
        self._by_reference: dict[str, list[MarshallProduct]] = {}
        for product in products:
            keys = {number_key(product.article)} | {
                number_key(reference.number) for reference in product.references
            }
            for key in keys:
                self._by_reference.setdefault(key, []).append(product)

    @classmethod
    def from_workbook(cls, path: Path) -> "MarshallIndex":
        return cls(brake_pad_rows(path))

    def lookup(self, query: str, *, max_products: int) -> tuple[list[str], list[Cross]]:
        products = self._by_reference.get(number_key(query), [])[:max_products]
        output: list[Cross] = []
        seen: set[tuple[str, str, str]] = set()
        for product in products:
            candidates = [
                MarshallReference("MARSHALL", product.article, KIND_AFTERMARKET),
                *product.references,
            ]
            for reference in candidates:
                key = reference.brand, number_key(reference.number), product.article
                if key in seen:
                    continue
                seen.add(key)
                output.append(Cross(
                    brand=reference.brand,
                    number=reference.number,
                    kind=reference.kind,
                    source="marshall",
                    source_product=product.article,
                    url=product.url or HOMEPAGE,
                ))
        return [product.article for product in products], output


class SQLiteMarshallIndex:
    """Read-only runtime view of a reviewed MARSHALL brake-pad snapshot."""

    def __init__(self, path: Path):
        self.path = path.resolve()

    def lookup(self, query: str, *, max_products: int) -> tuple[list[str], list[Cross]]:
        uri = f"file:{self.path.as_posix()}?mode=ro"
        with sqlite3.connect(uri, uri=True) as connection:
            rows = connection.execute(
                """
                SELECT DISTINCT product_article, product_url
                FROM matches
                WHERE query_key = ?
                ORDER BY product_article
                LIMIT ?
                """,
                (number_key(query), max_products),
            ).fetchall()
            if not rows:
                return [], []
            articles = [row[0] for row in rows]
            placeholders = ",".join("?" for _ in articles)
            references = connection.execute(
                f"""
                SELECT product_article, product_url, brand, number, kind
                FROM crosses
                WHERE product_article IN ({placeholders})
                ORDER BY product_article, kind DESC, brand, number_key
                """,
                articles,
            ).fetchall()
        products: dict[str, MarshallProduct] = {
            article: MarshallProduct(article, url, ()) for article, url in rows
        }
        grouped: dict[str, list[MarshallReference]] = {article: [] for article in products}
        for article, url, brand, number, kind in references:
            products.setdefault(article, MarshallProduct(article, url, ()))
            grouped[article].append(MarshallReference(brand, number, kind))
        return MarshallIndex([
            MarshallProduct(product.article, product.url, tuple(grouped[article]))
            for article, product in products.items()
        ]).lookup(query, max_products=max_products)


def build_sqlite_index(workbook_path: Path, output_path: Path) -> dict:
    """Compile one reviewed official workbook and atomically replace the index."""
    workbook_path = workbook_path.resolve()
    products = brake_pad_rows(workbook_path)
    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(output_path.suffix + ".tmp")
    temporary.unlink(missing_ok=True)
    try:
        with sqlite3.connect(temporary) as connection:
            connection.executescript("""
                PRAGMA journal_mode=OFF;
                PRAGMA synchronous=FULL;
                CREATE TABLE matches (
                    product_article TEXT NOT NULL,
                    product_url TEXT NOT NULL,
                    query_key TEXT NOT NULL
                );
                CREATE TABLE crosses (
                    product_article TEXT NOT NULL,
                    product_url TEXT NOT NULL,
                    number_key TEXT NOT NULL,
                    brand TEXT NOT NULL,
                    number TEXT NOT NULL,
                    kind TEXT NOT NULL
                );
                CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
            """)
            match_rows = []
            reference_rows = []
            for product in products:
                refs = [
                    MarshallReference("MARSHALL", product.article, KIND_AFTERMARKET),
                    *product.references,
                ]
                lookup_keys = {number_key(reference.number) for reference in refs}
                for lookup_key in lookup_keys:
                    match_rows.append((product.article, product.url, lookup_key))
                for reference in refs:
                    reference_rows.append((
                        product.article, product.url, number_key(reference.number),
                        reference.brand, reference.number, reference.kind,
                    ))
            connection.executemany("INSERT INTO matches VALUES (?, ?, ?)", match_rows)
            connection.executemany("INSERT INTO crosses VALUES (?, ?, ?, ?, ?, ?)", reference_rows)
            connection.executemany(
                "INSERT INTO metadata(key, value) VALUES (?, ?)", {
                    "schema_version": "1",
                    "products": str(len(products)),
                    "matches": str(len(match_rows)),
                    "references": str(len(reference_rows)),
                    "input_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
                    "source_url": CATALOGUE_URL,
                }.items(),
            )
            connection.executescript("""
                CREATE INDEX idx_matches_query
                    ON matches(query_key, product_article);
                CREATE INDEX idx_crosses_product
                    ON crosses(product_article, number_key);
                ANALYZE;
            """)
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)
    return {
        "path": str(output_path),
        "products": len(products),
        "matches": len(match_rows),
        "references": len(reference_rows),
        "size_bytes": output_path.stat().st_size,
        "sha256": hashlib.sha256(output_path.read_bytes()).hexdigest(),
        "input_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "source_url": CATALOGUE_URL,
    }


class MarshallSource(BaseSource):
    key = "marshall"
    title = "MARSHALL"
    homepage = HOMEPAGE
    verified = False
    groups = (BRAKE_PADS,)
    note = "Candidate W1: локальный индекс из публичного XLSX MARSHALL; выключен до отдельного gate."

    def __init__(self, settings, http_factory=None, pool=None, *, index=None):
        super().__init__(settings, http_factory, pool)
        self._index = index

    def _load_index(self) -> MarshallIndex | SQLiteMarshallIndex:
        if self._index is not None:
            return self._index
        configured = getattr(self.settings, "marshall_index_path", "")
        path = Path(configured) if configured else None
        if path is None or not path.is_file():
            raise FileNotFoundError("локальный индекс MARSHALL не настроен")
        self._index = SQLiteMarshallIndex(path)
        return self._index

    async def lookup(self, oe: str) -> SourceResult:
        started = self.timer()
        try:
            products, crosses = await asyncio.to_thread(
                self._load_index().lookup, oe,
                max_products=self.settings.max_products_per_oe,
            )
        except (OSError, ValueError, sqlite3.Error) as exc:
            return SourceResult(self.key, SourceStatus.ERROR, message=str(exc)[:300],
                                elapsed_ms=self.elapsed(started), url=HOMEPAGE)
        return SourceResult(
            self.key, SourceStatus.OK if products else SourceStatus.NOT_FOUND,
            crosses=crosses, products=products, elapsed_ms=self.elapsed(started),
            url=HOMEPAGE,
        )
