"""Сверить выдачу сервиса с эталонными списками из файла заказчика.

    python3 scripts/compare_with_reference.py [эталон.xlsx] [результат.xlsx]

Эталон — «2-й вариант» на первом листе клиентского файла: артикул и через
запятую номера, которые заказчик ждёт увидеть.
"""
import pathlib
import re
import sys
import xml.etree.ElementTree as ET
import zipfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from app.normalize import number_key

REFERENCE = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else pathlib.Path("/root/Тест - парсинг кроссов.xlsx")
RESULT = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else pathlib.Path("out/result.xlsx")

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
_CELL = re.compile(r"([A-Z]+)(\d+)")


def reference_lists(path: pathlib.Path) -> dict[str, list[str]]:
    """Пары «артикул -> список номеров» из клиентского файла.

    Колонки ищем по содержимому, а не по фиксированным координатам: файлы
    заказчика приходят с разной шапкой.
    """
    z = zipfile.ZipFile(path)
    shared = [
        "".join(t.text or "" for t in si.iter("{%s}t" % NS["m"]))
        for si in ET.fromstring(z.read("xl/sharedStrings.xml")).findall("m:si", NS)
    ]
    sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
    cells = {}
    for c in sheet.iter("{%s}c" % NS["m"]):
        v = c.find("m:v", NS)
        if v is None:
            continue
        cells[c.get("r")] = shared[int(v.text)] if c.get("t") == "s" else v.text

    out: dict[str, list[str]] = {}
    for ref, value in cells.items():
        if not isinstance(value, str) or value.count(",") < 3:
            continue
        col, row = _CELL.match(ref).groups()
        left = cells.get(f"{chr(ord(col[-1]) - 1)}{row}") if len(col) == 1 else None
        if not left or " " in str(left).strip():
            continue
        numbers = [x.strip() for x in value.split(",") if x.strip()]
        if len(numbers) > 3:
            out[str(left).strip()] = numbers
    return out


def collected(path: pathlib.Path) -> dict[str, set[str]]:
    ws = load_workbook(path)["Вариант 2"]
    header = [str(c.value or "") for c in ws[1]]
    sku_at = header.index("Наш артикул")
    numbers_at = next(i for i, h in enumerate(header) if h in ("Все кроссы", "Номера кроссов"))
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[sku_at]:
            continue
        out[row[sku_at]] = {number_key(x) for x in str(row[numbers_at] or "").split(",") if x.strip()}
    return out


def main() -> int:
    expected = reference_lists(REFERENCE)
    got = collected(RESULT)
    if not expected:
        print(f"В {REFERENCE.name} не нашлось эталонных списков")
        return 1

    total_expected = total_hit = 0
    for sku, numbers in expected.items():
        have = got.get(sku, set())
        hits = [n for n in numbers if number_key(n) in have]
        miss = [n for n in numbers if number_key(n) not in have]
        total_expected += len(numbers)
        total_hit += len(hits)
        print(f"{sku}: ожидалось {len(numbers):3d}, найдено {len(hits):3d} "
              f"({100 * len(hits) // len(numbers)}%), всего собрано {len(have)}")
        if miss:
            print(f"   не найдено: {', '.join(miss[:12])}{' …' if len(miss) > 12 else ''}")
    print(f"\nИТОГО покрытие эталона: {total_hit}/{total_expected} = "
          f"{100 * total_hit // total_expected}%")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
