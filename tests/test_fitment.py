from __future__ import annotations

import io
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app import main
from app.config import Settings
from app.fitment import FitmentService
from app.excel import build_fitment_workbook
from app.models import Base
from app.schemas import FitmentLookupRequest
from app.sources.trialli_fitment import TrialliFitmentSource
from app.sources.kyb_fitment import KybFitmentSource
from app.sources.torr_fitment import TorrFitmentSource


CARD = """
<html><head><title>Амортизатор задний AG 29125 купить недорого</title></head><body>
<h1>Амортизатор для LADA XRAY задний</h1>
<table class="props"><tr><td>Ось (сторона установки):</td><td>Сзади, слева/справа</td></tr>
<tr><td>Тип амортизатора:</td><td>Двухтрубный</td></tr></table>
<table class="js-table-car"><tr><th>Марка</th><th>Модель</th><th>Модификация</th>
<th>Код двигателя</th><th>Мощность, лс</th><th>Объём двигателя</th><th>Год выпуска</th></tr>
<tr><td>LADA</td><td>XRAY (GAB_)</td><td>1.6 Cross</td><td>H4M</td><td>113</td>
<td>1598</td><td>2019 - н.в.</td></tr>
<tr><td>SUZUKI</td><td>SX4</td><td>1.6 4x4</td><td>M16A</td><td>107</td>
<td>1586</td><td>2006 - 2013</td></tr></table>
<table class="js-table-car"><tr><th>Марка</th><th>Модель</th><th>Двигатель</th>
<th>Мощность, лс</th><th>Объём двигателя</th><th>Год</th></tr>
<tr><td>LADA</td><td>XRAY</td><td>H4M</td><td>113</td><td>1598</td><td>2019 - н.в.</td></tr></table>
</body></html>
"""


def test_trialli_fitment_parser_uses_full_table_and_keeps_period_semantics():
    result = TrialliFitmentSource.parse_page(CARD, source_url="https://trialli.ru/card/")
    assert result is not None
    assert result["number"] == "AG 29125"
    assert result["installation_position"] == "Сзади, слева/справа"
    assert len(result["applications"]) == 2  # reduced mobile duplicate is ignored
    assert result["applications"][0] == {
        "make": "LADA",
        "model": "XRAY (GAB_)",
        "modification": "1.6 Cross",
        "engine_code": "H4M",
        "power_hp": "113",
        "engine_cc": "1598",
        "raw_period": "2019 - н.в.",
        "year_from": 2019,
        "year_to": None,
        "end_is_open": True,
    }
    assert result["applications"][1]["year_to"] == 2013
    assert result["applications"][1]["end_is_open"] is False


@pytest.mark.asyncio
async def test_fitment_service_caches_complete_answer(tmp_path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'fitment.db'}")
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)

    service = FitmentService(Settings(_env_file=None), sessions)
    answer = TrialliFitmentSource.parse_page(CARD, source_url="https://trialli.ru/card/")
    service.sources["TRIALLI"].lookup = AsyncMock(return_value=answer)

    first = await service.lookup("TRIALLI", "AG 29125")
    second = await service.lookup("TRIALLI", "AG29125")
    assert first["cached"] is False
    assert second["cached"] is True
    assert len(second["applications"]) == 2
    service.sources["TRIALLI"].lookup.assert_awaited_once()
    await engine.dispose()


@pytest.mark.asyncio
async def test_fitment_endpoint_is_server_side_admin_only(monkeypatch):
    lookup = AsyncMock(return_value={"status": "ok", "applications": []})
    monkeypatch.setattr(main.fitment, "lookup", lookup)
    request = FitmentLookupRequest(
        brand="TRIALLI", number="AG 29125", group="shock_absorbers"
    )

    with pytest.raises(HTTPException) as denied:
        await main.fitment_lookup(request, tenant="another-user")
    assert denied.value.status_code == 403
    lookup.assert_not_awaited()

    assert await main.fitment_lookup(request, tenant="admin") == {
        "status": "ok", "applications": []
    }
    lookup.assert_awaited_once_with("TRIALLI", "AG 29125")


@pytest.mark.asyncio
async def test_fitment_endpoint_rejects_other_brands_and_groups(monkeypatch):
    lookup = AsyncMock()
    monkeypatch.setattr(main.fitment, "lookup", lookup)
    with pytest.raises(HTTPException) as brand_error:
        await main.fitment_lookup(
            FitmentLookupRequest(brand="SACHS", number="333754", group="shock_absorbers"),
            tenant="admin",
        )
    assert brand_error.value.status_code == 400
    with pytest.raises(HTTPException) as group_error:
        await main.fitment_lookup(
            FitmentLookupRequest(brand="TRIALLI", number="PF0806", group="brake_pads"),
            tenant="admin",
        )
    assert group_error.value.status_code == 400
    lookup.assert_not_awaited()


TORR_CARD = """
<div class="product-page__title">Shock absorber, Gas, LEFT/RIGHT, REAR Art: DH1270</div>
<table class="item-page__applicability"><thead><tr><th>MARKE</th><th>MODELL</th>
<th>BAUJAHR</th><th>MODIFIKATION</th><th>MOTOR</th><th>HUBRAUM</th><th>KW/PS</th><th>INFO</th>
</tr></thead><tbody><tr><td>KIA</td><td><span class="tooltip" data-tooltip="Rio II Sedan (JB)">Rio II Sed...</span></td>
<td>2005-03/-</td><td><span class="tooltip" data-tooltip="Rio 1.4 16V">Rio 1.4...</span></td>
<td>G4EE</td><td>1399</td><td>71/97</td><td>Standard</td></tr></tbody></table>
<table class="item-page__applicability"><tbody><tr><td>OEM 553101G200</td></tr></tbody></table>
"""


def test_torr_fitment_parser_uses_applicability_table_and_tooltips():
    result = TorrFitmentSource.parse_page(TORR_CARD, source_url="https://controltorr.de/catalog/DH1270")
    assert result is not None
    assert result["number"] == "DH1270"
    assert result["applications"] == [{
        "make": "KIA", "model": "Rio II Sedan (JB)", "modification": "Rio 1.4 16V",
        "engine_code": "G4EE", "engine_cc": "1399", "power_kw": "71", "power_hp": "97",
        "raw_period": "2005-03/-", "year_from": 2005, "year_to": None,
        "end_is_open": True, "info": "Standard",
    }]


def test_kyb_fitment_parser_preserves_open_ended_period():
    result = KybFitmentSource.parse_data({
        "partNumber": "333754",
        "application": [{"startMonth": 6, "startYear": 2006, "endMonth": 99,
                         "endYear": 9999, "firm": "Suzuki", "model": "SX4"}],
        "spec": {"part": "Амортизатор", "series": "Excel-G",
                 "type": "Двухтрубная газонаполненная стойка", "fr": "Передний",
                 "rl": "Левый", "fixType": "Стойка МакФерсон"},
    })
    assert result["status"] == "ok"
    assert result["applications"][0]["raw_period"] == "06.2006 — н.в."
    assert result["applications"][0]["year_to"] is None
    assert result["applications"][0]["end_is_open"] is True


def test_fitment_workbook_has_product_and_application_sheets():
    result = KybFitmentSource.parse_data({
        "partNumber": "333754",
        "application": [{"startMonth": 6, "startYear": 2006, "endMonth": 99,
                         "endYear": 9999, "firm": "Suzuki", "model": "SX4"}],
        "spec": {"part": "Амортизатор", "series": "Excel-G", "fr": "Передний"},
    })
    workbook = load_workbook(io.BytesIO(build_fitment_workbook(result)), data_only=True)
    assert workbook.sheetnames == ["Деталь", "Применяемость"]
    assert workbook["Применяемость"]["A2"].value == "KYB"
    assert workbook["Применяемость"]["C2"].value == "Suzuki"
    assert workbook["Применяемость"]["D2"].value == "SX4"
    workbook.close()


@pytest.mark.asyncio
async def test_fitment_export_is_server_side_admin_only(monkeypatch):
    lookup = AsyncMock(return_value=KybFitmentSource.parse_data({
        "partNumber": "333754", "application": [{"startYear": 2006, "endYear": 9999,
        "firm": "Suzuki", "model": "SX4"}], "spec": {"part": "Амортизатор"},
    }))
    monkeypatch.setattr(main.fitment, "lookup", lookup)
    request = FitmentLookupRequest(brand="KYB", number="333754", group="shock_absorbers")
    with pytest.raises(HTTPException) as denied:
        await main.fitment_export(request, tenant="rnd")
    assert denied.value.status_code == 403
    lookup.assert_not_awaited()
    response = await main.fitment_export(request, tenant="admin")
    assert response.media_type.endswith("spreadsheetml.sheet")
    assert response.body.startswith(b"PK")
    lookup.assert_awaited_once_with("KYB", "333754")
