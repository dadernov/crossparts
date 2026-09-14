import io

from openpyxl import load_workbook
import pytest

from app.main import lookup_export
from app.schemas import LookupExportCross, LookupExportRequest


@pytest.mark.asyncio
async def test_single_lookup_export_returns_canonical_excel_numbers():
    response = await lookup_export(LookupExportRequest(
        oe_number="58101-H5A25",
        group="brake_pads",
        crosses=[LookupExportCross(
            brand="Brembo", number="P 30-122 / A", kind="aftermarket", sources=["brembo"],
        )],
    ), tenant="test")

    workbook = load_workbook(io.BytesIO(response.body))
    first_variant = list(workbook["Вариант 1"].iter_rows(values_only=True))
    second_variant = list(workbook["Вариант 2"].iter_rows(values_only=True))
    assert first_variant[1][3] == "P30122A"
    assert second_variant[1][5] == "P30122A"
    assert response.headers["content-disposition"] == 'attachment; filename="crosses-58101H5A25.xlsx"'
