import pathlib, sys
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

import io
from openpyxl import Workbook, load_workbook

from app.excel import build_template, build_workbook, read_input


def test_oe_only_column():
    items = read_input(_sheet([["Номер ОЕ"], ["58101H5A25"]]))
    assert len(items) == 1
    assert items[0]['oe_number'] == '58101H5A25'
    assert items[0]['our_sku'] == ''


def test_template_has_no_sample_data_and_preserves_gaps():
    assert read_input(build_template()) == []
    workbook = load_workbook(io.BytesIO(build_template()))
    workbook.active.append(['A', 'Деталь', '', '111222'])
    workbook.active.append([])
    workbook.active.append(['B', 'Другая деталь', '', '333444'])
    buffer = io.BytesIO()
    workbook.save(buffer)
    items = read_input(buffer.getvalue())
    assert [item['oe_number'] for item in items] == ['111222', '333444']
    assert items[1]['part_name'] == 'Другая деталь'

SAMPLE = pathlib.Path("/root/Тест - парсинг кроссов.xlsx")


def _sheet(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_input_finds_header_row():
    data = _sheet([
        ["Наш Артикул", "Номер ОЕ для парсинга кроссов"],
        ["BPF159CG", "58101H5A25"],
        ["BPF394PG", "2L2Z2001BA"],
    ])
    assert [(i["our_sku"], i["oe_number"], i["group"]) for i in read_input(data)] == [
        ("BPF159CG", "58101H5A25", None),
        ("BPF394PG", "2L2Z2001BA", None),
    ]


def test_read_input_stops_at_blank_row_before_other_blocks():
    data = _sheet([
        ["Наш Артикул", "Номер ОЕ"],
        ["A1", "111222"],
        [None, None],
        ["Сайты для парсинга", None],
        ["Brembo", "https://example.com"],
    ])
    items = read_input(data)
    assert [(i["our_sku"], i["oe_number"]) for i in items] == [("A1", "111222")]


def test_read_input_on_the_real_client_file():
    if not SAMPLE.exists():
        return
    items = read_input(SAMPLE.read_bytes())
    assert len(items) == 14
    assert (items[0]["our_sku"], items[0]["oe_number"]) == ("BPF159CG", "58101H5A25")


def test_build_workbook_layout():
    blob = build_workbook([{
        "our_sku": "BPF159CG",
        "oe_number": "58101H5A25",
        "group": "brake_pads",
        "status": "ok",
        "crosses": [
            {"brand": "HYUNDAI", "number": "58101-H5A25", "kind": "oem", "sources": ["sbparts"]},
            {"brand": "NIBK", "number": "PN0537", "kind": "aftermarket", "sources": ["sbparts"]},
            {"brand": "BRANNOR", "number": "XV40", "kind": "aftermarket", "sources": ["brannor"]},
        ],
        "source_reports": [{"source": "sbparts", "status": "ok", "crosses": 2}],
    }])
    wb = load_workbook(io.BytesIO(blob))
    assert wb.sheetnames == ["Вариант 1", "Вариант 2"]
    v1 = list(wb["Вариант 1"].iter_rows(values_only=True))
    assert v1[0] == ("Номер ОЕ (запрос)", "Бренд аналога", "Номер аналога",
                     "Раздел", "Источники")
    assert v1[1] == ("58101H5A25", "HYUNDAI", "58101H5A25", "OEM", "sbparts")
    v2 = list(wb["Вариант 2"].iter_rows(values_only=True))
    assert v2[0] == ("Наш артикул", "Товарная группа", "Номер ОЕ (запрос)",
                     "Кол-во", "ОЕМ/Афтермаркет", "Все кроссы")
    assert v2[1] == ("BPF159CG", "Тормозные колодки", "58101H5A25", 1,
                     "ОЕМ", "58101H5A25")
    assert v2[2] == ("BPF159CG", "Тормозные колодки", "58101H5A25", 1,
                     "АФТЕРМАРКЕТ", "PN0537")
    assert "XV40" not in str(v1) + str(v2)


def test_export_strips_all_special_characters_from_cross_numbers():
    blob = build_workbook([{
        "our_sku": "A1", "oe_number": "12345", "group": None,
        "crosses": [
            {"brand": "TEST", "number": "ab-12 / 34.5", "kind": "aftermarket", "sources": ["test"]},
            {"brand": "OTHER", "number": "AB 12-34/5", "kind": "aftermarket", "sources": ["test"]},
        ],
    }])
    workbook = load_workbook(io.BytesIO(blob))
    variant_one = list(workbook["Вариант 1"].iter_rows(values_only=True))
    variant_two = list(workbook["Вариант 2"].iter_rows(values_only=True))
    assert variant_one[1][2] == "AB12345"
    assert variant_one[2][2] == "AB12345"
    assert variant_two[1][3] == 1
    assert variant_two[1][5] == "AB12345"


def test_unrecognised_group_is_shown_as_sent():
    """Незнакомую группу не выдумываем — показываем как прислали."""
    blob = build_workbook([{
        "our_sku": "X1", "oe_number": "111222",
        "group": None, "group_raw": "Свечи зажигания", "status": "ok",
        "crosses": [{"brand": "NGK", "number": "BKR6E", "kind": "aftermarket",
                     "sources": ["x"]}],
        "source_reports": [],
    }])
    wb = load_workbook(io.BytesIO(blob))
    assert list(wb["Вариант 2"].iter_rows(values_only=True))[1][1] == "Свечи зажигания"
